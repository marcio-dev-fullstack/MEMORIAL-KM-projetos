import os
import json
from docx import Document
from datetime import datetime

def gerar_memorial_de_json(caminho_json, pasta_saida):
    with open(caminho_json, 'r', encoding='utf-8') as f:
        dados = json.load(f)

    doc = Document()
    doc.add_heading('MEMORIAL DESCRITIVO E AVALIAÇÃO TÉCNICA', 0)
    data_atual = datetime.now().strftime('%d de %B de %Y')
    doc.add_paragraph(f"Conceição do Araguaia - PA, {data_atual}")
    
    doc.add_heading('1. Descrição do Imóvel', level=1)
    doc.add_paragraph(dados['descricao'])
    
    doc.add_heading('2. Orçamento Sintético', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = 'Item', 'Descrição', 'Valor (R$)'
    
    for item in dados['orcamento']:
        row = table.add_row().cells
        row[0].text = item['item']
        row[1].text = item['desc']
        row[2].text = f"{item['valor']:,.2f}"
    
    doc.add_heading('3. Responsabilidade Técnica', level=1)
    doc.add_paragraph(f"ART Vinculada: {dados.get('num_art', 'Não informada')}")
    doc.add_paragraph("\n\n__________________________________________")
    doc.add_paragraph(f"{dados['engenheiro']}\nEngenheiro Civil\nCREA-PA: {dados['crea']}")
    
    if not os.path.exists(pasta_saida):
        os.makedirs(pasta_saida)
        
    nome_arquivo = f"Memorial_{dados['projeto'].replace(' ', '_')}.docx"
    doc.save(os.path.join(pasta_saida, nome_arquivo))
    print(f"Gerado: {nome_arquivo}")

# Processar todos os arquivos na pasta 'dados'
pasta_dados = 'dados'
pasta_saida = 'output'

for arquivo in os.listdir(pasta_dados):
    if arquivo.endswith('.json'):
        gerar_memorial_de_json(os.path.join(pasta_dados, arquivo), pasta_saida)