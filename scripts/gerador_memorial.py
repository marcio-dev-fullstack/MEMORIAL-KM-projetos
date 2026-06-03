import json
from docx import Document
from docx.shared import Pt
from datetime import datetime

def gerar_memorial_de_json(arquivo_json):
    # Carregar dados
    with open(arquivo_json, 'r', encoding='utf-8') as f:
        dados = json.load(f)

    doc = Document()
    
    # Cabeçalho com Data
    doc.add_heading('MEMORIAL DESCRITIVO E AVALIAÇÃO TÉCNICA', 0)
    data_atual = datetime.now().strftime('%d de %B de %Y')
    doc.add_paragraph(f"Conceição do Araguaia - PA, {data_atual}")
    
    # Descrição do Imóvel
    doc.add_heading('1. Descrição do Imóvel', level=1)
    doc.add_paragraph(dados['descricao'])
    
    # Orçamento Sintético
    doc.add_heading('2. Orçamento Sintético', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = 'Item', 'Descrição', 'Valor (R$)'
    
    for item in dados['orcamento']:
        row = table.add_row().cells
        row[0].text = item['item']
        row[1].text = item['desc']
        row[2].text = f"{item['valor']:,.2f}"
    
    # Seção Normativa e Assinatura
    doc.add_heading('3. Responsabilidade Técnica', level=1)
    doc.add_paragraph("Este documento cumpre as diretrizes da norma ABNT NBR 14653.")
    doc.add_paragraph("\n\n__________________________________________")
    doc.add_paragraph(f"{dados['engenheiro']}\nEngenheiro Civil\nCREA-PA: {dados['crea']}")
    
    nome_arquivo = f"Memorial_{dados['projeto'].replace(' ', '_')}.docx"
    doc.save(nome_arquivo)
    print(f"Sucesso: {nome_arquivo} gerado com sucesso!")

# Executar
gerar_memorial_de_json('dados_imovel.json')