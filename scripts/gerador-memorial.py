import os
import json
import csv
from docx import Document
from datetime import datetime
from docx2pdf import convert

def registrar_log(dados, vt_total):
    """Auditoria de emissões."""
    # Garante que a pasta logs existe
    if not os.path.exists('logs'):
        os.makedirs('logs')
        
    file_exists = os.path.isfile('logs/registro_emissoes.csv')
    with open('logs/registro_emissoes.csv', 'a', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow(['Data', 'Projeto', 'ART', 'Valor'])
        writer.writerow([datetime.now().strftime('%d/%m/%Y'), dados['projeto'], dados.get('num_art', 'N/A'), vt_total])

def validar_lei_municipal(dados):
    """Validação de geometria legal."""
    recuo_f = dados.get('recuo_frontal', 0)
    if recuo_f < 5.0:
        return "AVISO: Recuo frontal em DESCONFORMIDADE com a lei municipal (mínimo 5m)."
    return "O imóvel encontra-se em CONFORMIDADE com a legislação municipal vigente."

def calcular_valores(d):
    """Cálculo financeiro NBR 14653."""
    vt = d['terreno']['area_m2'] * d['terreno']['valor_unitario_m2']
    cr = d['edificacao']['area_m2'] * d['edificacao']['cub_atual_m2']
    depreciacao = (d['edificacao']['idade_aparente_anos'] / d['edificacao']['vida_util_total_anos']) * cr
    valor_edificacao = (cr - depreciacao) * d['edificacao']['estado_conservacao']
    return vt, valor_edificacao, vt + valor_edificacao

def gerar_memorial(dados, base_dir):
    # Cálculos
    vt, ve, vt_total = calcular_valores(dados)
    
    # Validação Legal
    validacao = validar_lei_municipal(dados)
    
    # Docx
    doc = Document()
    doc.add_heading('LAUDO TÉCNICO DE AVALIAÇÃO', 0)
    doc.add_paragraph(f"Data: {datetime.now().strftime('%d/%m/%Y')}")
    doc.add_heading('Parecer Legal', level=1)
    doc.add_paragraph(validacao)
    
    doc.add_heading('Avaliação Econômica (NBR 14653)', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    for label, valor in [('Valor Terreno', vt), ('Valor Edificação', ve), ('VALOR TOTAL', vt_total)]:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = f"R$ {valor:,.2f}"
    
    # Salvar
    if not os.path.exists(os.path.join(base_dir, 'output')):
        os.makedirs(os.path.join(base_dir, 'output'))
        
    caminho_docx = os.path.join(base_dir, 'output', f"Laudo_{dados['projeto'].replace(' ', '_')}.docx")
    doc.save(caminho_docx)
    
    # Exportação PDF (Nota: certifique-se de que o Word esteja instalado no Windows para o docx2pdf funcionar)
    try:
        convert(caminho_docx)
    except Exception as e:
        print(f"Erro ao converter PDF: {e}")
        
    registrar_log(dados, vt_total)